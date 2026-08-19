#!/usr/bin/env python3
"""Build a simple LMSA PLUS exec ops hub workbook for Google Sheets."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "ops/LMSA-PLUS-Exec-Hub.xlsx"

NAVY = "003057"
GOLD = "B3A369"
CREAM = "F8F7F2"
WHITE = "FFFFFF"
INK = "102033"

OFFICERS = [
    ("President", "Stacy Lomeli", "Vision, LMSA National/Southeast, GT relationships, approvals, board alignment"),
    ("Vice President / Operations", "Amanda Ubera-Corona", "Task tracker, deadlines, meeting follow-ups, calendar, member systems"),
    ("Secretary", "Open", "Meeting notes, records, attendance, shared docs, board communications"),
    ("Treasurer / Finance", "Open", "Fundraising, dues, reimbursements, purchases, GT funding"),
    ("Outreach / Partnerships", "Ashley Diaz Duenas", "Clinics, schools, nonprofits, medical schools, GT offices, other LMSA chapters"),
    ("Service / Community Health", "Open", "Volunteer events, health fairs, screenings, Hispanic-family outreach"),
    ("Events / Programming", "Open", "Speakers, panels, workshops, GBMs, networking, Medical Spanish"),
    ("Marketing / Communications", "Open", "Instagram, graphics, announcements, website, promotion, recaps"),
]

STATUS_LIST = '"Not started,In progress,Blocked,Done,Cancelled"'
PRIORITY_LIST = '"High,Medium,Low"'
TYPE_LIST = '"Chapter event,Campus deadline,Partner meeting,Internal deadline,Other"'
OWNER_LIST = (
    '"President,Vice President / Operations,Secretary,Treasurer / Finance,'
    'Outreach / Partnerships,Service / Community Health,Events / Programming,'
    'Marketing / Communications,Whole board"'
)

thin = Border(
    left=Side(style="thin", color="D0D5DD"),
    right=Side(style="thin", color="D0D5DD"),
    top=Side(style="thin", color="D0D5DD"),
    bottom=Side(style="thin", color="D0D5DD"),
)
header_fill = PatternFill("solid", fgColor=NAVY)
header_font = Font(name="Calibri", bold=True, color=WHITE, size=11)
title_font = Font(name="Calibri", bold=True, color=NAVY, size=16)
gold_font = Font(name="Calibri", bold=True, color=GOLD, size=11)
body_font = Font(name="Calibri", color=INK, size=11)
cream_fill = PatternFill("solid", fgColor=CREAM)
wrap = Alignment(wrap_text=True, vertical="top")


def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def prep_rows(ws, start_row, cols, n=40):
    for r in range(start_row, start_row + n):
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = thin
            cell.font = body_font
            cell.alignment = wrap
            if r % 2 == 0:
                cell.fill = cream_fill


def add_validation(ws, formula, cells):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(cells)


wb = Workbook()

# ---------- HOW TO USE ----------
ws = wb.active
ws.title = "START HERE"
ws["A1"] = "LMSA PLUS at Georgia Tech — Exec Hub (simple)"
ws["A1"].font = title_font
ws["A3"] = "What this is"
ws["A3"].font = gold_font
ws["A4"] = (
    "One shared workbook for the whole board. Use Google Sheets so everyone can edit at the same time. "
    "Keep updates short. If a task needs an owner, put it on Exec Tasks and assign a role."
)
ws["A4"].alignment = wrap

ws["A6"] = "How to use (2 minutes)"
ws["A6"].font = gold_font
steps = [
    "1. Upload this file to Google Drive → Open with Google Sheets → Share with the board (Editor).",
    "2. Calendar tab = dates everyone needs to see (events, deadlines, partner meetings).",
    "3. Exec Tasks tab = the main to-do list. Anyone can ADD a task and ASSIGN an officer.",
    "4. Each officer tab = that person's own checklist. Move finished items to Done.",
    "5. Update Status every meeting (Amanda / VP leads this).",
    "6. Optional: also create one shared Google Calendar and paste event links in the Calendar tab Notes.",
]
for i, s in enumerate(steps, start=7):
    ws[f"A{i}"] = s
    ws[f"A{i}"].alignment = wrap

ws["A14"] = "Do NOT use this for"
ws["A14"].font = gold_font
ws["A15"] = (
    "Private student data, passwords, unapproved personal phone numbers, or unfinished partnerships "
    "you are not ready to share with the full board."
)
ws["A15"].alignment = wrap

ws["A17"] = "Current officers"
ws["A17"].font = gold_font
ws["A18"] = "Role"
ws["B18"] = "Name"
ws["C18"] = "Focus"
style_header(ws, 18, 3)
for i, (role, name, focus) in enumerate(OFFICERS, start=19):
    ws[f"A{i}"] = role
    ws[f"B{i}"] = name
    ws[f"C{i}"] = focus
    for col in ("A", "B", "C"):
        ws[f"{col}{i}"].font = body_font
        ws[f"{col}{i}"].alignment = wrap
        ws[f"{col}{i}"].border = thin
set_widths(ws, [32, 24, 70])
ws.row_dimensions[4].height = 45
ws.row_dimensions[15].height = 40

# ---------- CALENDAR ----------
cal = wb.create_sheet("Calendar")
cal["A1"] = "Main Calendar — add / remove dates here"
cal["A1"].font = title_font
cal["A2"] = "Anyone can add a row. Keep one row = one date. Put the owner so people know who owns it."
cal["A2"].font = body_font
headers = ["Date", "Time (optional)", "What", "Type", "Owner", "Status", "Notes / link"]
for i, h in enumerate(headers, start=1):
    cal.cell(row=4, column=i, value=h)
style_header(cal, 4, 7)
prep_rows(cal, 5, 7, 50)
add_validation(cal, TYPE_LIST, "D5:D60")
add_validation(cal, OWNER_LIST, "E5:E60")
add_validation(cal, STATUS_LIST, "F5:F60")
# sample rows
samples = [
    ("2026-08-06", "", "Board chair applications open", "Internal deadline", "President", "In progress", ""),
    ("2026-08-27", "5:00–6:00 PM", "Pre-Health Fall Kickoff — Klaus Atrium", "Chapter event", "Events / Programming", "Not started", "Confirmed"),
    ("2026-09-01", "11:00 AM–1:00 PM", "Student Org Fair Day 1 — Tech Green", "Chapter event", "Outreach / Partnerships", "Not started", "Confirmed"),
    ("2026-09-02", "11:00 AM–1:00 PM", "Student Org Fair Day 2 — Tech Green", "Chapter event", "Outreach / Partnerships", "Not started", "Confirmed"),
]
for r, row in enumerate(samples, start=5):
    for c, val in enumerate(row, start=1):
        cal.cell(row=r, column=c, value=val)
set_widths(cal, [12, 16, 42, 18, 26, 14, 28])
cal.freeze_panes = "A5"

# ---------- EXEC TASKS ----------
tasks = wb.create_sheet("Exec Tasks")
tasks["A1"] = "Exec Tasks — main board to-do list"
tasks["A1"].font = title_font
tasks["A2"] = (
    "This is the hub. To add work: insert a new row → write the task → choose Owner → set Due date → Status = Not started. "
    "Amanda (VP / Operations) reviews this every meeting."
)
tasks["A2"].font = body_font
tasks["A2"].alignment = wrap
headers = ["Task", "Owner", "Due date", "Priority", "Status", "Notes"]
for i, h in enumerate(headers, start=1):
    tasks.cell(row=4, column=i, value=h)
style_header(tasks, 4, 6)
prep_rows(tasks, 5, 6, 60)
add_validation(tasks, OWNER_LIST, "B5:B70")
add_validation(tasks, PRIORITY_LIST, "D5:D70")
add_validation(tasks, STATUS_LIST, "E5:E70")
starter = [
    ("Publish board chair Google Form and add link to website", "President", "2026-08-06", "High", "Not started", ""),
    ("Create shared Google Calendar and invite all officers", "Vice President / Operations", "2026-08-10", "High", "Not started", "Paste calendar link in Notes"),
    ("Confirm Interest Meeting date/time/location", "Events / Programming", "2026-09-01", "High", "Not started", ""),
    ("Draft Org Fair table checklist (brochures, QR, goodies)", "Outreach / Partnerships", "2026-08-25", "Medium", "Not started", ""),
    ("Post Kickoff Instagram reminder", "Marketing / Communications", "2026-08-24", "Medium", "Not started", ""),
]
for r, row in enumerate(starter, start=5):
    for c, val in enumerate(row, start=1):
        tasks.cell(row=r, column=c, value=val)
set_widths(tasks, [55, 28, 12, 10, 14, 36])
tasks.row_dimensions[2].height = 40
tasks.freeze_panes = "A5"

# ---------- OFFICER TABS ----------
for role, name, focus in OFFICERS:
    short = {
        "President": "01 President",
        "Vice President / Operations": "02 VP Operations",
        "Secretary": "03 Secretary",
        "Treasurer / Finance": "04 Treasurer Finance",
        "Outreach / Partnerships": "05 Outreach",
        "Service / Community Health": "06 Service",
        "Events / Programming": "07 Events",
        "Marketing / Communications": "08 Marketing",
    }[role]
    sheet = wb.create_sheet(short)
    sheet["A1"] = role
    sheet["A1"].font = title_font
    sheet["A2"] = f"Officer: {name}"
    sheet["A2"].font = gold_font
    sheet["A3"] = f"Focus: {focus}"
    sheet["A3"].font = body_font
    sheet["A3"].alignment = wrap
    sheet["A5"] = "My checklist — add your own tasks here. Move Done items to the bottom or mark Status = Done."
    sheet["A5"].font = body_font
    headers = ["Task", "Due date", "Status", "Blocked by?", "Notes"]
    for i, h in enumerate(headers, start=1):
        sheet.cell(row=7, column=i, value=h)
    style_header(sheet, 7, 5)
    prep_rows(sheet, 8, 5, 35)
    add_validation(sheet, STATUS_LIST, "C8:C50")
    set_widths(sheet, [50, 12, 14, 22, 36])
    sheet.freeze_panes = "A8"
    sheet.row_dimensions[3].height = 35

# ---------- MEETING LOG (tiny) ----------
meet = wb.create_sheet("Meeting Log")
meet["A1"] = "Meeting Log — short notes only"
meet["A1"].font = title_font
meet["A2"] = "Secretary fills this after each meeting. Keep it to decisions + owners."
headers = ["Date", "What we decided", "Who owns follow-up", "Due", "Done?"]
for i, h in enumerate(headers, start=1):
    meet.cell(row=4, column=i, value=h)
style_header(meet, 4, 5)
prep_rows(meet, 5, 5, 30)
add_validation(meet, OWNER_LIST, "C5:C40")
add_validation(meet, STATUS_LIST, "E5:E40")
set_widths(meet, [12, 50, 28, 12, 12])
meet.freeze_panes = "A5"

import os
os.makedirs("ops", exist_ok=True)
wb.save(OUT)
print(f"Wrote {OUT}")
